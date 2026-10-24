"""图内文字检测与提取；直转通道（pdf2docx）输出的 docx 后处理。"""

import hashlib
import io

from PIL import Image

MIN_CHARS = 6      # 判定「图中有文字」的最少字符数
MIN_CONF = 0.65    # 平均置信度阈值
MARK = "图内文字："  # 可编辑文本前缀标记


def _ocr_result_texts(engine, np_arr):
    """跑一遍 OCR，返回 [(text, conf), ...]。"""
    results = engine.predict(np_arr)
    lines = []
    for res in results:
        d = dict(res) if not isinstance(res, dict) else res
        texts = d.get("rec_texts") or []
        scores = d.get("rec_scores") or [1.0] * len(texts)
        for i, t in enumerate(texts):
            t = str(t).strip()
            if t:
                conf = float(scores[i]) if i < len(scores) else 1.0
                lines.append((t, conf))
    return lines


def detect_image_text(img, engine):
    """检测图片中是否含文字；有则返回可编辑文本，否则返回 None。

    img: PIL.Image 或图片字节。
    """
    import numpy as np

    if isinstance(img, (bytes, bytearray)):
        try:
            img = Image.open(io.BytesIO(img))
        except Exception:
            return None
    try:
        img = img.convert("RGB")
    except Exception:
        return None
    if img.width < 60 or img.height < 40:  # 过小的装饰图跳过
        return None
    lines = _ocr_result_texts(engine, np.asarray(img))
    if not lines:
        return None
    total_chars = sum(len(t) for t, _ in lines)
    avg_conf = sum(c for _, c in lines) / len(lines)
    if total_chars >= MIN_CHARS and avg_conf >= MIN_CONF:
        return "".join(t for t, _ in lines)
    return None


def count_docx_images(docx_path):
    """统计 docx 中嵌入的图片数量（纯 XML 扫描，不需要 OCR 引擎）。

    用于在加载 OCR 模型前判断是否存在需要检测的图片，
    避免无图文档白白付出模型加载开销。
    """
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(docx_path)
    total = 0
    paras = list(doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                paras.extend(cell.paragraphs)
    for para in paras:
        total += len(para._p.findall(".//" + qn("a:blip")))
    return total


def enrich_direct_docx(docx_path, engine, log=print):
    """在 pdf2docx 输出的 docx 中，把图片里的文字提取为可编辑段落插在图后。

    返回命中的图片段落数。
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document(docx_path)
    paras = list(doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                paras.extend(cell.paragraphs)

    cache, hits = {}, 0
    gray = RGBColor(0x59, 0x59, 0x59)
    for para in paras:
        blips = para._p.findall(".//" + qn("a:blip"))
        texts = []
        for blip in blips:
            rid = blip.get(qn("r:embed"))
            if not rid:
                continue
            try:
                blob = doc.part.related_parts[rid].blob
            except Exception:
                continue
            key = hashlib.sha1(blob).hexdigest()
            if key not in cache:
                cache[key] = detect_image_text(blob, engine)
            if cache[key]:
                texts.append(cache[key])
        if not texts:
            continue
        combined = "\n".join(dict.fromkeys(texts))
        new_p = doc.add_paragraph()
        r1 = new_p.add_run(MARK)
        r1.bold = True
        r1.italic = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = gray
        r2 = new_p.add_run(combined)
        r2.italic = True
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = gray
        para._p.addnext(new_p._p)  # 紧跟在含图段落之后
        hits += 1

    if hits:
        doc.save(docx_path)
        log(f"检测到 {hits} 处图片含文字，已提取为可编辑段落")
    return hits
