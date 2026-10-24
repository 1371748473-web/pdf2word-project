"""重建管线：版面分析 / 文本层提取 / OCR → 段落重组 → docx 输出。

块（block）格式：
    ("para", text, pt, bold[, style])   style 可选 "gray"（图内文字标注）
    ("image", data_bytes, ext[, width_ratio])   width_ratio 为图宽占页宽比例
    ("table", grid)   grid=[[{"text":str,"colspan":int},...],...]
"""

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

MAX_IMG_WIDTH_CM = 15.8   # A4 去掉页边距后的可用宽度近似
MIN_IMG_PIXELS = 80       # 过滤装饰性小图的边长阈值（像素）

_ocr_engine = None


def get_ocr_engine():
    """惰性初始化 PaddleOCR（进程内单例，避免重复加载模型）。

    设备由 envsetup._detect_gpu() 通过 nvidia-smi 检测，
    首次创建引擎时调用 activate_gpu_device() 设置 paddle.set_device("gpu")。
    """
    global _ocr_engine
    if _ocr_engine is None:
        from .envsetup import activate_gpu_device, patch_paddlex_deps
        from paddleocr import PaddleOCR

        # paddleocr 已把 paddlex 正常导入，此时打补丁不会触发重复初始化
        patch_paddlex_deps()
        activate_gpu_device()
        _ocr_engine = PaddleOCR(
            use_doc_orientation_classify=False,
            use_doc_unwarping=False,
            use_textline_orientation=False,
            # 禁止开启 oneDNN：paddlepaddle 3.3.0 + PP-OCRv6 下 det 模型在
            # oneDNN 指令集缺少 DoubleAttribute 支持，predictor.run() 直接抛
            # NotImplementedError（见 experiments/onednn/README.md，2026-09-19 实测）。
            # 升级 paddle 后可用 experiments/onednn/bench_onednn.py 复测。
            enable_mkldnn=False,
            lang="ch",
        )
    return _ocr_engine


# ---------------------------------------------------------------- OCR 部分

def _parse_ocr_results(results):
    """把 PaddleOCR 结果解析为行级列表（含坐标）。"""
    import numpy as np

    lines = []
    for res in results:
        d = dict(res) if not isinstance(res, dict) else res
        texts = d.get("rec_texts") or []
        scores = d.get("rec_scores") or [1.0] * len(texts)
        polys = d.get("rec_polys")
        if polys is None:
            polys = d.get("dt_polys")
        for i, t in enumerate(texts):
            t = str(t).strip()
            if not t:
                continue
            poly = polys[i] if polys is not None and i < len(polys) else None
            if poly is not None:
                arr = np.asarray(poly, dtype=float)
                x0, y0 = arr.min(axis=0)
                x1, y1 = arr.max(axis=0)
            else:
                x0, y0, x1, y1 = 0.0, 0.0, 1.0, 1.0
            lines.append(
                {
                    "text": t,
                    "conf": float(scores[i]) if i < len(scores) else 1.0,
                    "x0": float(x0),
                    "y0": float(y0),
                    "x1": float(x1),
                    "y1": float(y1),
                    "cy": (float(y0) + float(y1)) / 2.0,
                    "h": max(float(y1) - float(y0), 1.0),
                }
            )
    return lines


def ocr_lines_arr(np_arr):
    """对 numpy 图像 OCR，返回行级结果列表。"""
    engine = get_ocr_engine()
    # PaddleOCR 3.x 仅接受 numpy.ndarray / str 路径输入
    return _parse_ocr_results(engine.predict(np_arr))


def ocr_lines_arr_batch(np_arrs):
    """批量 OCR：输入 numpy 数组列表，返回每张图的行级结果列表。

    优先尝试批量推理以减少 GPU 调用开销；若批量接口不可用则逐张回退。
    """
    if not np_arrs:
        return []
    engine = get_ocr_engine()
    try:
        results = engine.predict(list(np_arrs))
        if not isinstance(results, list):
            results = [results]
        if len(results) == len(np_arrs):
            return [_parse_ocr_results([r]) for r in results]
    except Exception:
        pass
    return [_parse_ocr_results(engine.predict(a)) for a in np_arrs]


def ocr_lines(pil_image):
    import numpy as np

    return ocr_lines_arr(np.asarray(pil_image))


def group_lines(lines):
    """按阅读顺序排序并聚合成段落。"""
    if not lines:
        return []
    lines = sorted(lines, key=lambda L: (L["cy"], L["x0"]))
    heights = sorted(L["h"] for L in lines)
    med_h = heights[len(heights) // 2]
    paras, cur = [], [lines[0]]
    for L in lines[1:]:
        prev = cur[-1]
        gap = L["y0"] - prev["y1"]
        if gap > 0.6 * med_h:  # 垂直间距过大 → 新段落
            paras.append(cur)
            cur = [L]
        else:
            cur.append(L)
    paras.append(cur)
    return paras


def render_page(page, dpi=200):
    """把 PDF 页渲染为 PIL Image。"""
    import numpy as np
    from PIL import Image

    pix = page.get_pixmap(dpi=dpi)
    arr = np.frombuffer(pix.samples, dtype=np.uint8)
    if pix.n >= 4:
        arr = arr.reshape(pix.height, pix.width, pix.n)[:, :, :3]
    else:
        arr = arr.reshape(pix.height, pix.width, pix.n)
    return Image.fromarray(arr, "RGB")


def _crop_padded(img, x0, y0, x1, y1, pad=4):
    W, H = img.size
    x0 = max(0, int(x0) - pad)
    y0 = max(0, int(y0) - pad)
    x1 = min(W, int(x1) + pad)
    y1 = min(H, int(y1) + pad)
    if x1 - x0 < 10 or y1 - y0 < 10:
        return None
    return img.crop((x0, y0, x1, y1)), (x1 - x0) / W


# ------------------------------------------------------------ 扫描页重建

def _emit_text_from_lines(lines, label, dpi, blocks):
    """将已 OCR 的行列表按段落聚合并追加到 blocks。"""
    is_title = label is not None and "title" in label.lower()
    for group in group_lines(lines):
        text = "".join(L["text"] for L in group)
        if is_title:
            pt, bold = (16.0, True) if label == "doc_title" else (14.0, True)
        else:
            avg_h = sum(L["h"] for L in group) / len(group)
            pt = max(9.0, min(22.0, avg_h * 72.0 / dpi))
            bold = False
        blocks.append(("para", text, round(pt, 1), bold))


def scan_page_blocks(page, dpi=200, extract_image_text=True, log=None):
    """扫描页 → 块列表。优先版面分析分区处理，失败时退化为整页 OCR。"""
    import numpy as np
    from PIL import Image

    from . import imagetext, layout

    log = log or (lambda s: None)
    img = render_page(page, dpi=dpi)
    arr = np.asarray(img)
    engine = get_ocr_engine()
    blocks = []

    # 版面分析使用低 DPI 加速推理，坐标后续缩放回原 DPI
    layout_dpi = min(dpi, 150)
    layout_scale = layout_dpi / dpi
    if layout_scale < 1.0:
        lw = max(1, int(img.width * layout_scale))
        lh = max(1, int(img.height * layout_scale))
        layout_arr = np.asarray(img.resize((lw, lh), Image.LANCZOS))
    else:
        layout_arr = arr

    regions = layout.detect_regions(layout_arr)
    if not regions:
        log("    版面模型不可用，退化为整页 OCR")
        for group in group_lines(ocr_lines_arr(arr)):
            text = "".join(L["text"] for L in group)
            avg_h = sum(L["h"] for L in group) / len(group)
            pt = max(9.0, min(22.0, avg_h * 72.0 / dpi))
            blocks.append(("para", text, round(pt, 1), False))
        return blocks

    # 缩放区域坐标回原 DPI
    if layout_scale < 1.0:
        inv = 1.0 / layout_scale
        for r in regions:
            r["x0"] *= inv
            r["y0"] *= inv
            r["x1"] *= inv
            r["y1"] *= inv

    tasks = []          # (y0, kind, payload, y1, x0)
    covered = []        # 已被区域覆盖的纵向区间
    for r in regions:
        kind = layout.label_kind(r["label"])
        if kind == "skip":
            continue
        tasks.append((r["y0"], kind, r, r["y1"], r["x0"]))
        covered.append((r["y0"], r["y1"]))

    # 兜底：版面模型可能漏检个别文本行，对未覆盖的横向条带补跑 OCR
    H = img.size[1]
    min_band = max(24, int(H * 0.015))
    covered.sort()
    cur = 0
    for y0, y1 in covered:
        if y0 - cur > min_band:
            tasks.append((cur, "band", (cur, y0), y0, 0.0))
        cur = max(cur, y1)
    if H - cur > min_band:
        tasks.append((cur, "band", (cur, H), H, 0.0))

    # 排序：先按行聚类（y 区间重叠超过较短者 40% 视为同一行），行内按 x0 从左到右
    tasks.sort(key=lambda t: t[0])
    ordered, row, row_start, row_end = [], None, None, None
    for t in tasks:
        y0, y1 = t[0], t[3]
        if row is None:
            row, row_start, row_end = [t], y0, y1
            continue
        overlap = min(row_end, y1) - max(row_start, y0)
        h_min = min(row_end - row_start, y1 - y0)
        if h_min > 0 and overlap >= 0.4 * h_min:
            row.append(t)
            row_start = min(row_start, y0)
            row_end = max(row_end, y1)
        else:
            ordered.extend(sorted(row, key=lambda x: x[4]))
            row, row_start, row_end = [t], y0, y1
    if row:
        ordered.extend(sorted(row, key=lambda x: x[4]))

    from . import table as tablemod

    # ---- 预收集所有需要 OCR 的裁剪区域，批量推理 ----
    task_crops = {}   # task_idx → (crop_pil, width_ratio) 或 None
    ocr_inputs = []   # [(task_idx, np_arr), ...]

    for idx, (_, kind, payload, _, _) in enumerate(ordered):
        if kind == "band":
            b0, b1 = payload
            crop = img.crop((0, max(0, b0 - 2), img.size[0], min(H, b1 + 2)))
            task_crops[idx] = (crop, None)
            ocr_inputs.append((idx, np.asarray(crop)))
        else:
            r = payload
            cropped = _crop_padded(img, r["x0"], r["y0"], r["x1"], r["y1"])
            if not cropped:
                task_crops[idx] = None
                continue
            crop, width_ratio = cropped
            task_crops[idx] = (crop, width_ratio)
            if kind in ("text", "table"):
                ocr_inputs.append((idx, np.asarray(crop)))

    # 批量 OCR：一次推理处理所有文本/表格/条带裁剪图
    ocr_results = {}  # task_idx → lines list
    if ocr_inputs:
        batch_lines = ocr_lines_arr_batch([c for _, c in ocr_inputs])
        for (task_idx, _), lines in zip(ocr_inputs, batch_lines):
            ocr_results[task_idx] = lines

    # ---- 按阅读顺序处理结果，输出 blocks ----
    for idx, (_, kind, payload, _, _) in enumerate(ordered):
        crop_info = task_crops.get(idx)
        if crop_info is None:
            continue
        crop, width_ratio = crop_info

        if kind == "band":
            _emit_text_from_lines(ocr_results.get(idx, []), None, dpi, blocks)
        elif kind == "text":
            _emit_text_from_lines(ocr_results.get(idx, []), payload["label"], dpi, blocks)
        elif kind == "table":
            lines = ocr_results.get(idx, [])
            grid = tablemod.reconstruct_table(crop, lines, log)
            if grid:
                blocks.append(("table", grid))
            else:
                buf = io.BytesIO()
                crop.save(buf, "PNG")
                blocks.append(("image", buf.getvalue(), "png", round(width_ratio, 3)))
        else:  # image
            buf = io.BytesIO()
            crop.save(buf, "PNG")
            blocks.append(("image", buf.getvalue(), "png", round(width_ratio, 3)))
            if extract_image_text:
                t = imagetext.detect_image_text(crop, engine)
                if t:
                    blocks.append(("para", imagetext.MARK + t, 10.5, False, "gray"))
    return blocks


# ------------------------------------------------------------ 文本层部分

def extract_text_blocks(page, engine=None, extract_image_text=True):
    """文字页 → 块列表。利用文本层，保真度高于 OCR；图片做图内文字检测。"""
    from . import imagetext

    d = page.get_text("dict")
    blocks = []
    for block in d.get("blocks", []):
        btype = block.get("type", 0)
        if btype == 0:
            parts, size_sum, n, bold = [], 0.0, 0, False
            for ln in block.get("lines", []):
                for sp in ln.get("spans", []):
                    t = sp.get("text", "")
                    if not t.strip():
                        continue
                    parts.append(t)
                    size_sum += sp.get("size", 12.0)
                    n += 1
                    font = sp.get("font", "")
                    if "Bold" in font or "bold" in font:
                        bold = True
                parts.append("\n")
            text = "".join(parts).strip()
            if text:
                size = size_sum / n if n else 12.0
                blocks.append(("para", text.replace("\n", ""), round(size, 1), bold))
        elif btype == 1:
            w = block.get("width", 0)
            h = block.get("height", 0)
            data = block.get("image")
            if not data or w < MIN_IMG_PIXELS or h < MIN_IMG_PIXELS:
                continue
            blocks.append(("image", data, block.get("ext", "png"), None))
            if extract_image_text and engine is not None:
                t = imagetext.detect_image_text(data, engine)
                if t:
                    blocks.append(("para", imagetext.MARK + t, 10.5, False, "gray"))
    return blocks


# ------------------------------------------------------------ DOCX 组装

def _setup_styles(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(12)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "宋体")


def _add_table_block(doc, grid):
    """把识别出的表格网格渲染为带边框的 Word 表格。"""
    if not grid:
        return
    ncols = max(sum(c.get("colspan", 1) for c in row) for row in grid)
    ncols = max(ncols, 1)
    try:
        tbl = doc.add_table(rows=len(grid), cols=ncols)
        tbl.style = "Table Grid"
        tbl.autofit = True
    except Exception:
        return
    for i, row in enumerate(grid):
        col = 0
        for cell in row:
            if col >= ncols:
                break
            cs = max(1, int(cell.get("colspan", 1)))
            cs = min(cs, ncols - col)
            try:
                tc = tbl.cell(i, col)
                if cs > 1:
                    tc = tc.merge(tbl.cell(i, col + cs - 1))
                para = tc.paragraphs[0]
                r = para.add_run(str(cell.get("text", "")))
                r.font.size = Pt(9)
                rpr = r._element.get_or_add_rPr()
                rfonts = rpr.get_or_add_rFonts()
                rfonts.set(qn("w:eastAsia"), "宋体")
            except Exception:
                pass
            col += cs


def build_docx(out_path, pages_blocks):
    """pages_blocks: [ [block, ...], ... ]。"""
    doc = Document()
    _setup_styles(doc)
    gray = RGBColor(0x59, 0x59, 0x59)
    for idx, blocks in enumerate(pages_blocks, 1):
        if idx > 1:
            doc.add_page_break()
        mark = doc.add_paragraph()
        mark.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = mark.add_run(f"— 第 {idx} 页 —")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        for blk in blocks:
            if blk[0] == "para":
                text, pt, bold = blk[1], blk[2], blk[3]
                style = blk[4] if len(blk) > 4 else None
                para = doc.add_paragraph()
                r = para.add_run(text)
                r.font.size = Pt(pt)
                r.bold = bold
                if style == "gray":
                    r.italic = True
                    r.font.color.rgb = gray
            elif blk[0] == "table":
                _add_table_block(doc, blk[1])
            elif blk[0] == "image":
                data = blk[1]
                ratio = blk[3] if len(blk) > 3 else None
                if ratio:
                    width = Cm(min(MAX_IMG_WIDTH_CM, max(3.0, MAX_IMG_WIDTH_CM * ratio)))
                else:
                    width = Cm(MAX_IMG_WIDTH_CM)
                para = doc.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                try:
                    para.add_run().add_picture(io.BytesIO(data), width=width)
                except Exception as e:  # 个别异常图片跳过，不阻断整体
                    para.add_run(f"[图片插入失败: {e}]")
    doc.save(out_path)
